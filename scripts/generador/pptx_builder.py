from pptx import Presentation
from pptx.util import Pt
from io import BytesIO
from pathlib import Path
import pandas as pd
from config import TEMPLATE_PPTX, OLE_POSITIONS


def _reemplazar_texto(slide, nombre_shape: str, texto: str):
    """Reemplaza texto buscando por nombre exacto del shape."""
    for shape in slide.shapes:
        if shape.name == nombre_shape and shape.has_text_frame:
            tf = shape.text_frame
            for para in tf.paragraphs:
                for run in para.runs:
                    run.text = ""
            if tf.paragraphs and tf.paragraphs[0].runs:
                tf.paragraphs[0].runs[0].text = texto
            elif tf.paragraphs:
                tf.paragraphs[0].text = texto
            return


def _shape_por_nombre(slide, nombre: str):
    """Retorna el shape con ese nombre o None."""
    for shape in slide.shapes:
        if shape.name == nombre:
            return shape
    return None


def _insertar_imagen(slide, img_bytes: bytes, pos: dict):
    stream = BytesIO(img_bytes)
    slide.shapes.add_picture(
        stream,
        left=pos["left"], top=pos["top"],
        width=pos["width"], height=pos["height"],
    )


def generar_pptx(
    nombre_fondo:    str,
    periodo_str:     str,
    comentario_pm:   str,
    img_evolucion:   bytes,
    img_tabla_rent:  bytes,
    img_composicion: bytes,
    img_tabla_comp:  bytes,
    info_fondo:      dict,
    out_path:        Path,
) -> Path:
    prs    = Presentation(str(TEMPLATE_PPTX))
    slide0 = prs.slides[0]
    slide1 = prs.slides[1]

    # ── Slide 0: nombre del fondo ─────────────────────────────────
    nombre_corto = nombre_fondo.replace("FIP VANTRUST ", "")
    _reemplazar_texto(slide0, "CuadroTexto 6", nombre_corto)

    # ── Slide 0: comentario PM ────────────────────────────────────
    _reemplazar_texto(slide0, "CuadroTexto 13", comentario_pm)

    # ── Slide 0: información general (CuadroTexto 30) ─────────────
    shape_vals = _shape_por_nombre(slide0, "CuadroTexto 30")
    if shape_vals and shape_vals.has_text_frame:
        tf = shape_vals.text_frame
        valores_info = [
            info_fondo.get("administradora", "Vantrust Gestion Patrimonial S.A."),
            info_fondo.get("rut",            "76,637,334-8"),
            info_fondo.get("moneda",         "CLP"),
            info_fondo.get("tipo",           "Fondo de Inversión Privado"),
            info_fondo.get("fecha_inicio",   ""),
            info_fondo.get("benchmark",      "Índice Cámara Promedio (ICP)"),
            "Sí",
            info_fondo.get("plazo_rescate",  "A más tardar 15 días corridos"),
        ]
        for i, para in enumerate(tf.paragraphs):
            if i < len(valores_info):
                if para.runs:
                    para.runs[0].text = valores_info[i]
                else:
                    para.text = valores_info[i]

    # ── Slide 0: imágenes (reemplazan los OLE) ────────────────────
    _insertar_imagen(slide0, img_tabla_rent,  OLE_POSITIONS["tabla_rentabilidad"])
    _insertar_imagen(slide0, img_evolucion,   OLE_POSITIONS["grafico_evolucion"])
    _insertar_imagen(slide0, img_composicion, OLE_POSITIONS["grafico_composicion"])
    _insertar_imagen(slide0, img_tabla_comp,  OLE_POSITIONS["tabla_comparacion"])

    # ── Slide 1: nombre del fondo ─────────────────────────────────
    _reemplazar_texto(slide1, "CuadroTexto 15", nombre_corto)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path
