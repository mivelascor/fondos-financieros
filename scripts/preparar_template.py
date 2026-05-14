"""
preparar_template.py — Limpia el PPTX original eliminando los OLE vinculados
y guarda el template listo para usar por el generador.

INSTRUCCIONES:
1. Copia tu archivo PPTX original a scripts/templates/folleto_template.pptx
2. Corre este script UNA SOLA VEZ desde la carpeta scripts/:
       python preparar_template.py
3. Se genera templates/folleto_template_clean.pptx
4. Sube ese archivo al repo (scripts/templates/folleto_template_clean.pptx)
5. No vuelvas a correr esto a menos que cambies el diseño del folleto.
"""
from pptx import Presentation
from pathlib import Path

INPUT  = Path("templates/folleto_template.pptx")
OUTPUT = Path("templates/folleto_template_clean.pptx")

if not INPUT.exists():
    print(f"ERROR: No se encontró {INPUT}")
    print("Copia tu PPTX original a scripts/templates/folleto_template.pptx y vuelve a correr.")
    exit(1)

prs    = Presentation(str(INPUT))
slide0 = prs.slides[0]

print("=== SHAPES EN SLIDE 0 ===")
for i, shape in enumerate(slide0.shapes):
    tipo = shape.shape_type
    nombre = shape.name
    txt = ""
    if shape.has_text_frame:
        txt = shape.text_frame.text[:50].replace("\n", "|")
    extra = ""
    if shape.has_table:
        extra = f" [TABLE {len(shape.table.rows)}x{len(shape.table.columns)}]"
    if tipo == 10:
        extra = f" [OLE — left={shape.left} top={shape.top} w={shape.width} h={shape.height}]"
    print(f"  [{i:2d}] type={tipo:2d}  name={nombre!r:35s} {txt[:40]}{extra}")

print()

# Eliminar los OLE de la slide 0
ole_shapes = [s for s in slide0.shapes if s.shape_type == 10]
print(f"Eliminando {len(ole_shapes)} objetos OLE vinculados:")
for s in ole_shapes:
    print(f"  - {s.name}: left={s.left} top={s.top} w={s.width} h={s.height}")
    s._element.getparent().remove(s._element)

print()
prs.save(str(OUTPUT))
print(f"✓ Template limpio guardado en: {OUTPUT}")
print()
print("IMPORTANTE: Copia las posiciones OLE que aparecen arriba a config.py")
print("en la variable OLE_POSITIONS si difieren de los valores actuales.")
