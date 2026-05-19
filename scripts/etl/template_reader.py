"""
etl/template_reader.py

Lee los valores calculados por LibreOffice directamente de la hoja
'rentabilidad' y 'Datos ICP (2)' del template Excel actualizado.

Retorna los mismos campos que espera pptx_builder.py:
  resumen, historico, grafico, nombre_fip, acum_label
"""
import openpyxl
from pathlib import Path

TEMPLATES_DIR = Path(__file__).parent.parent.parent / "inputs" / "templates"

TEMPLATE_MAP = {
    "FIP VANTRUST LIQUIDEZ ACTIVA":        "TEMPLATE_FONDO_LIQUIDEZ_ACTIVA.xlsx",
    "FIP VANTRUST LIQUIDEZ ALTO APORTE":   "TEMPLATE_FONDO_ALTO_APORTE.xlsx",
    "FIP VANTRUST LIQUIDEZ ALTO CAPITAL":  "TEMPLATE_FONDO_ALTO_CAPITAL.xlsx",
    "FIP VANTRUST LIQUIDEZ ALTO MONTO":    "TEMPLATE_FONDO_LIQUIDEZ_ALTO_MONTO.xlsx",
    "FIP VANTRUST LIQUIDEZ CAJA":          "TEMPLATE_FONDO_LIQUIDEZ_CAJA.xlsx",
    "FIP VANTRUST LIQUIDEZ CONTINUA":      "TEMPLATE_FONDO_LIQUIDEZ_CONTINUA.xlsx",
    "FIP VANTRUST LIQUIDEZ CORRIENTE":     "TEMPLATE_FONDO_LIQUIDEZ_CORRIENTE.xlsx",
    "FIP VANTRUST LIQUIDEZ CORTO PLAZO":   "TEMPLATE_FONDO_LIQUIDEZ_CORTO_PLAZO.xlsx",
    "FIP VANTRUST LIQUIDEZ DISPONIBLE I":  "TEMPLATE_FONDO_LIQUIDEZ_Disponible_I.xlsx",
    "FIP VANTRUST LIQUIDEZ DOLAR":         "TEMPLATE_FONDO_LIQUIDEZ_DOLAR.xlsx",
    "FIP VANTRUST LIQUIDEZ DOLAR CAJA":    "TEMPLATE_FONDO_LIQUIDEZ_DOLAR_CAJA.xlsx",
    "FIP VANTRUST LIQUIDEZ EFECTIVO":      "TEMPLATE_FONDO_LIQUIDEZ_EFECTIVO.xlsx",
    "FIP VANTRUST LIQUIDEZ FLEXIBLE":      "TEMPLATE_FONDO_LIQUIDEZ_FLEXIBLE.xlsx",
    "FIP VANTRUST LIQUIDEZ I":             "TEMPLATE_FONDO_LIQUIDEZ_UNO.xlsx",
    "FIP VANTRUST LIQUIDEZ LOCAL":         "TEMPLATE_FONDO_LIQUIDEZ_LOCAL.xlsx",
    "FIP VANTRUST LIQUIDEZ MONETARIO I":   "TEMPLATE_FONDO_LIQUIDEZ_Monetario_I.xlsx",
    "FIP VANTRUST LIQUIDEZ PERMANENTE":    "TEMPLATE_FONDO_LIQUIDEZ_Permanente.xlsx",
    "FIP VANTRUST LIQUIDEZ PLUS":          "TEMPLATE_FONDO_LIQUIDEZ_PLUS.xlsx",
    "FIP VANTRUST LIQUIDEZ PRESENTE":      "TEMPLATE_FONDO_LIQUIDEZ_Presente.xlsx",
    "FIP VANTRUST LIQUIDEZ RECURRENTE":    "TEMPLATE_FONDO_LIQUIDEZ.xlsx",
    "FIP VANTRUST LIQUIDEZ RENDIMIENTO":   "TEMPLATE_FONDO_LIQUIDEZ_RENDIMIENTO.xlsx",
    "FIP VANTRUST LIQUIDEZ RESERVA DOLAR": "TEMPLATE_FONDO_LIQUIDEZ_RESERVA_DOLAR.xlsx",
    "FIP VANTRUST LIQUIDEZ SENCILLO":      "TEMPLATE_FONDO_LIQUIDEZ_SENCILLO.xlsx",
    "FIP VANTRUST LIQUIDEZ TEMPORAL":      "TEMPLATE_FONDO_USD_MONEY_MARKET.xlsx",
}


def _f(v):
    try:
        return float(v) if v is not None else None
    except (TypeError, ValueError):
        return None


def leer_datos_template(nombre_fondo: str) -> dict:
    """
    Lee el template Excel (ya recalculado por LibreOffice) y retorna:
    {
      acum_label : str
      nombre_fip : str
      resumen    : [{nombre, m, t, s, a, ac, es_icp, es_comp, es_fip}]
      historico  : [{año, filas:[{nombre, meses:[12], total}]}]
      grafico    : {labels:[], icp:[], comp:[], fip:[]}
    }
    """
    archivo = TEMPLATE_MAP.get(nombre_fondo)
    if not archivo:
        raise KeyError(f"Sin template para: {nombre_fondo}")
    ruta = TEMPLATES_DIR / archivo
    if not ruta.exists():
        raise FileNotFoundError(f"Template no encontrado: {ruta}")

    # data_only=True lee los valores calculados por LibreOffice
    wb = openpyxl.load_workbook(str(ruta), read_only=True, data_only=True)

    # ── Hoja rentabilidad ─────────────────────────────────────────────────
    ws   = wb["rentabilidad"]
    rows = list(ws.iter_rows(values_only=True))

    # Fila 0: headers. Col 23 = label acum
    acum_raw   = rows[0][23] if rows[0] and len(rows[0]) > 23 else "Acum."
    acum_label = str(acum_raw).strip().replace("\n", " ") if acum_raw else "Acum."

    # Filas 1-3: resumen
    resumen    = []
    nombre_fip = nombre_fondo.replace("FIP VANTRUST ", "FIP ")

    for i in range(1, 4):
        if i >= len(rows): break
        row = rows[i]
        if len(row) < 24: continue
        nombre = str(row[18]).strip() if row[18] else None
        if not nombre or _f(row[19]) is None: continue
        es_icp  = "icp"    in nombre.lower() or "benchmark" in nombre.lower()
        es_comp = "compet" in nombre.lower()
        es_fip  = not es_icp and not es_comp
        if es_fip: nombre_fip = nombre
        resumen.append({
            "nombre":  nombre,
            "m":  _f(row[19]), "t": _f(row[20]),
            "s":  _f(row[21]), "a": _f(row[22]),
            "ac": _f(row[23]),
            "es_icp": es_icp, "es_comp": es_comp, "es_fip": es_fip,
        })

    # Filas 5+: histórico
    hist_map = {}
    año_cur  = None
    for row in rows[5:]:
        if len(row) < 17: continue
        if row[2] is not None:
            try: año_cur = int(row[2])
            except (TypeError, ValueError): continue
        if año_cur is None: continue
        serie = str(row[3]).strip() if row[3] else None
        if not serie: continue
        meses = [_f(row[j]) if j < len(row) else None for j in range(4, 16)]
        total = _f(row[16]) if len(row) > 16 else None
        if not any(v is not None and v != 0 for v in meses):
            continue
        if año_cur not in hist_map:
            hist_map[año_cur] = []
        hist_map[año_cur].append({"nombre": serie, "meses": meses, "total": total})

    historico = [{"año": a, "filas": hist_map[a]}
                 for a in sorted(hist_map.keys()) if hist_map[a]]

    # ── Hoja Datos ICP (2): niveles para gráfico ──────────────────────────
    grafico = {"labels": [], "icp": [], "comp": [], "fip": []}
    if "Datos ICP (2)" in wb.sheetnames:
        ws2  = wb["Datos ICP (2)"]
        rows2 = list(ws2.iter_rows(values_only=True))

        raw = []
        for row in rows2[1:]:
            f    = row[0]
            icp  = _f(row[1]) if len(row) > 1  else None  # col B
            fip  = _f(row[7]) if len(row) > 7  else None  # col H (nivel con divid.)
            comp = _f(row[10]) if len(row) > 10 else None  # col K
            if not (f and hasattr(f, "year")): continue
            if not any([icp, fip, comp]): continue
            raw.append({"fecha": f, "icp": icp, "fip": fip, "comp": comp})

        # Encontrar inicio del FIP
        primera_fip = next((r for r in raw if r["fip"] and r["fip"] > 0), None)
        if primera_fip:
            import pandas as pd
            f0 = pd.Timestamp(primera_fip["fecha"])
            puntos = [r for r in raw if pd.Timestamp(r["fecha"]) >= f0]
            v0_icp  = primera_fip.get("icp")  or 1
            v0_comp = primera_fip.get("comp") or 1
            v0_fip  = primera_fip.get("fip")  or 1
            for pt in puntos:
                import datetime
                f = pt["fecha"]
                if hasattr(f, "strftime"):
                    lbl = f.strftime("%b %Y")
                else:
                    lbl = str(f)[:7]
                grafico["labels"].append(lbl)
                grafico["icp"].append(round(pt["icp"]/v0_icp*100, 2) if pt["icp"] and v0_icp else None)
                grafico["comp"].append(round(pt["comp"]/v0_comp*100, 2) if pt["comp"] and v0_comp else None)
                grafico["fip"].append(round(pt["fip"]/v0_fip*100, 2) if pt["fip"] and v0_fip else None)

    wb.close()
    return {
        "acum_label": acum_label,
        "nombre_fip": nombre_fip,
        "resumen":    resumen,
        "historico":  historico,
        "grafico":    grafico,
    }
